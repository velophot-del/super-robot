"""课程 XLSX 解析和逐行校验。"""
from io import BytesIO
from openpyxl import load_workbook
import re
from zipfile import ZIP_DEFLATED, ZipFile
from app.utils.csv_parser import CSV_TEMPLATES, _convert_field_value, _normalize_course_schedule_fields
from app.services.course_library import normalize_course_classification


def _load_workbook_for_values(content: bytes):
    try:
        return load_workbook(BytesIO(content),read_only=True,data_only=True)
    except TypeError as exc:
        # Some Excel-compatible exporters write invalid fill styles. Import only
        # reads cell values, so retry after removing style metadata from the package.
        if 'openpyxl.styles.fills.Fill' not in str(exc):
            raise
    source=BytesIO(content);cleaned=BytesIO()
    with ZipFile(source) as original, ZipFile(cleaned,'w',ZIP_DEFLATED) as output:
        if sum(entry.file_size for entry in original.infolist())>50*1024*1024:
            raise ValueError('Excel文件解压后超过50MB')
        for entry in original.infolist():
            name=entry.filename
            if name=='xl/styles.xml':
                continue
            data=original.read(name)
            if name=='xl/_rels/workbook.xml.rels':
                data=re.sub(rb'<Relationship\b(?=[^>]*\bType="[^"]*/styles")[^>]*/>',b'',data)
            elif name=='[Content_Types].xml':
                data=re.sub(rb'<Override\b(?=[^>]*\bPartName="/xl/styles\.xml")[^>]*/>',b'',data)
            elif name.startswith('xl/worksheets/') and name.endswith('.xml'):
                data=re.sub(rb'(<c\b[^>]*?)\s+s="\d+"',rb'\1',data)
            output.writestr(entry,data)
    cleaned.seek(0)
    return load_workbook(cleaned,read_only=True,data_only=True)


def parse_course_workbook(content: bytes):
    workbook=_load_workbook_for_values(content)
    fields=CSV_TEMPLATES['courses']['required_fields']+CSV_TEMPLATES['courses']['optional_fields']
    aliases={name:field for field,name in CSV_TEMPLATES['courses']['field_names'].items()}
    aliases.update({field:field for field in fields})
    aliases.update({
        '学期教学次数':'semester_sessions', '每周上课次数':'weekly_sessions',
        '每次课时':'hours_per_session', '每次学时':'hours_per_session',
        'semester_sessions':'semester_sessions', 'weekly_sessions':'weekly_sessions',
        'hours_per_session':'hours_per_session', '总学时':'total_hours', '计划总学时':'total_hours',
        '适用专业':'department', '专业名称':'department', '专业代码':'program_code',
        '专业方向':'direction_name', '开课单位':'offering_unit_code', '教研室编码':'offering_unit_code',
    })
    rows=[];errors=[]
    for sheet in workbook.worksheets:
        values=list(sheet.iter_rows(values_only=True))
        header_index=None;header_map={}
        for index,raw in enumerate(values[:15]):
            mapped={}
            for column,value in enumerate(raw):
                key=str(value or '').strip()
                if key in aliases:mapped[column]=aliases[key]
            mapped_fields=set(mapped.values())
            if ({'name','semester','total_hours'}.issubset(mapped_fields)
                    or {'name','semester','semester_sessions'}.issubset(mapped_fields)):
                header_index=index;header_map=mapped;break
        if header_index is None:
            if any(any(value not in (None,'') for value in row) for row in values):
                errors.append({'sheet':sheet.title,'row':None,'message':'前15行未找到课程名称、上课学期及总课时（或旧版学期教学次数）表头'})
            continue
        for index,raw in enumerate(values[header_index+1:],start=header_index+2):
            if not any(value not in (None,'') for value in raw):continue
            row={}
            for column,field in header_map.items():
                if column<len(raw):row[field]=_convert_field_value(field,raw[column])
            _normalize_course_schedule_fields(row)
            row['_source_sheet']=sheet.title;row['_source_row']=index
            rows.append(row)
    workbook.close()
    return rows,errors


def parse_academic_catalog_workbook(content: bytes):
    """Parse the college's two-sheet major/direction workbook without trusting its styles."""
    workbook=_load_workbook_for_values(content)
    details=None;summary=None
    for sheet in workbook.worksheets:
        values=list(sheet.iter_rows(values_only=True))
        for index,row in enumerate(values[:15]):
            labels={str(value or '').strip() for value in row}
            if {'专业名称','专业代码','专业方向','对应教研室'}<=labels:
                details=(sheet.title,values,index,row)
            if {'专业名称','专业代码','方向数'}<=labels:
                summary=(sheet.title,values,index,row)
    if not details:
        workbook.close()
        raise ValueError('未找到包含“专业名称、专业代码、专业方向、对应教研室”的明细表头')
    detail_name,detail_values,detail_header,detail_row=details
    detail_cols={str(value or '').strip():index for index,value in enumerate(detail_row)}
    programs={};errors=[]
    for row_index,raw in enumerate(detail_values[detail_header+1:],start=detail_header+2):
        if not any(value not in (None,'') for value in raw):
            continue
        get=lambda label: raw[detail_cols[label]] if label in detail_cols and detail_cols[label]<len(raw) else None
        code=str(get('专业代码') or '').strip().removesuffix('.0')
        name=str(get('专业名称') or '').strip()
        if not code or not name:
            errors.append(f'{detail_name}第{row_index}行缺少专业代码或专业名称')
            continue
        direction=str(get('专业方向') or '').strip()
        if not direction:
            errors.append(f'{detail_name}第{row_index}行缺少专业方向/招生口径')
            continue
        try:
            duration=normalize_program_duration(get('学制'))
        except ValueError as exc:
            errors.append(f'{detail_name}第{row_index}行：{exc}')
            continue
        entry_type='整专业招生' if direction.startswith('整专业招生') else '专业方向'
        degree=str(get('授予学位') or '').strip()
        construction=str(get('专业建设级别') or '').strip()
        notes=str(get('备注') or '').strip()
        record=programs.setdefault(code,{'code':code,'name':name,'education_level':'本科',
            'degree_category':degree or None,'duration_semesters':duration,
            'construction_level':construction or None,'source_notes':None,'directions':[],
            '_notes':[],'_source_rows':[]})
        record['_source_rows'].append(row_index)
        if record['name']!=name:
            errors.append(f'专业代码{code}对应多个名称：{record["name"]} / {name}')
        if record['degree_category'] and degree and record['degree_category']!=degree:
            errors.append(f'专业代码{code}对应多个授予学位口径')
        if record['duration_semesters']!=duration:
            errors.append(f'专业代码{code}对应多个学制口径')
        if construction and construction not in [x for x in (record['construction_level'] or '').split('；') if x]:
            record['construction_level']='；'.join(filter(None,[record['construction_level'],construction]))
        if notes and notes not in record['_notes']:
            record['_notes'].append(notes)
        unit_name=str(get('对应教研室') or '').strip()
        record['directions'].append({'name':direction,'entry_type':entry_type,'offering_unit_name':unit_name or None,
            'notes':notes or None,'source_order':len(record['directions'])+1,'source_row':row_index})
    for record in programs.values():
        record['source_notes']='；'.join(record.pop('_notes')) or None
        record.pop('_source_rows',None)

    warnings=[]
    if summary:
        _,summary_values,summary_header,summary_row=summary
        summary_cols={str(value or '').strip():index for index,value in enumerate(summary_row)}
        detail_count={code:len(record['directions']) for code,record in programs.items()}
        for raw in summary_values[summary_header+1:]:
            if not any(value not in (None,'') for value in raw):
                continue
            code=str(raw[summary_cols['专业代码']] if summary_cols['专业代码']<len(raw) else '').strip().removesuffix('.0')
            if code not in programs:
                continue
            reported=raw[summary_cols['方向数']] if summary_cols['方向数']<len(raw) else None
            try:
                reported=int(reported)
            except (TypeError,ValueError):
                continue
            if reported!=detail_count[code]:
                warnings.append(f'{programs[code]["name"]}（{code}）汇总表方向数为{reported}，明细表有{detail_count[code]}条方向/招生口径记录')
            summary_name=str(raw[summary_cols['专业名称']] or '').strip() if summary_cols['专业名称']<len(raw) else ''
            if summary_name and summary_name!=programs[code]['name']:
                warnings.append(f'{code}在汇总表名称为“{summary_name}”，明细表名称为“{programs[code]["name"]}”')
    workbook.close()
    return list(programs.values()),errors,warnings


def normalize_program_duration(value):
    text=str(value or '').strip()
    match=re.search(r'(\d+|[一二三四五六七八九十])\s*年',text)
    if not match:
        raise ValueError(f'无法识别学制“{text}”')
    years=match.group(1)
    years=int(years) if years.isdigit() else {'一':1,'二':2,'三':3,'四':4,'五':5,'六':6,'七':7,'八':8,'九':9,'十':10}[years]
    if not 1<=years<=10:
        raise ValueError(f'学制“{text}”超出有效范围')
    return years*2


def course_row_issues(row, db):
    issues=[]
    for field,label in [('name','课程名称'),('semester','上课学期'),('total_hours','总课时')]:
        if row.get(field) in (None,''):
            issues.append(f'{label}不能为空')
    try:
        if not 1<=int(row.get('semester'))<=12:issues.append('上课学期须为1至12')
    except (TypeError,ValueError):issues.append('上课学期须为1至12')
    try:
        if int(row.get('total_hours'))<1:issues.append('总课时须大于0')
    except (TypeError,ValueError):issues.append('总课时须为整数')
    weekdays=row.get('weekdays') or []
    if weekdays and (not isinstance(weekdays,list) or any(not isinstance(day,int) or not 1<=day<=7 for day in weekdays)):
        issues.append('每周上课日须为周一至周日')
    if row.get('cohort') and not str(row['cohort']).strip().removesuffix('级').isdigit():
        issues.append('适用年级请填写年份，如2026')
    if row.get('teaching_weeks') not in (None, ''):
        try:
            if not 1<=int(row['teaching_weeks'])<=20:issues.append('授课周数须为1至20')
        except (TypeError,ValueError):issues.append('授课周数须为整数')
    from app.models.teacher import Teacher
    from app.models.planning import OfferingUnit, AcademicProgram, AcademicDirection
    code=row.get('leader_teacher_code')
    teacher=db.query(Teacher).filter_by(teacher_code=str(code).strip()).first() if code else None
    if code and not teacher:issues.append(f'课程负责人工号“{code}”不存在')
    unit_code=row.get('offering_unit_code')
    unit=db.query(OfferingUnit).filter_by(code=str(unit_code).strip()).first() if unit_code else None
    if unit_code and not unit:issues.append(f'开课单位编码“{unit_code}”不存在')
    program_code=str(row.get('program_code') or '').strip()
    program=db.query(AcademicProgram).filter_by(code=program_code).first() if program_code else None
    program_name=str(row.get('department') or '').strip()
    if program_code and not program:issues.append(f'专业编码“{program_code}”不存在')
    elif program_name and program and program.name!=program_name:issues.append('专业名称与专业编码不一致')
    elif program_name and not program:
        program=db.query(AcademicProgram).filter_by(name=program_name).first()
        if not program and db.query(AcademicProgram.id).first():issues.append(f'专业“{program_name}”不在管理员专业目录中')
    if program:
        row['department']=program.name
    direction_name=str(row.get('direction_name') or '').strip()
    direction=None
    if direction_name and program:
        direction=db.query(AcademicDirection).filter_by(program_id=program.id,name=direction_name,entry_type='专业方向').first()
        if not direction:issues.append(f'专业“{program.name}”下未找到方向“{direction_name}”')
    row['_resolved_program_id']=program.id if program else None
    row['_resolved_direction_id']=direction.id if direction else None
    row['_resolved_unit_id']=direction.offering_unit_id if direction and direction.offering_unit_id else (unit.id if unit else None)
    if direction and direction.offering_unit_id and unit and direction.offering_unit_id!=unit.id:
        issues.append('开课教研室与所选专业方向对应的教研室不一致')
    return issues,teacher.id if teacher else None,unit.id if unit else None


def normalize_course_row(row, teacher_id=None, unit_id=None):
    from app.models.course import Course
    allowed={column.name for column in Course.__table__.columns}
    result={key:value for key,value in row.items() if key in allowed and key not in {'id','created_at','updated_at'}}
    result['course_type'],result['course_nature']=normalize_course_classification(
        result.get('course_type') or '',result.get('course_nature') or '')
    if teacher_id:result['leader_teacher_id']=teacher_id
    if row.get('_resolved_unit_id'):result['offering_unit_id']=row['_resolved_unit_id']
    elif unit_id:result['offering_unit_id']=unit_id
    if row.get('_resolved_program_id'):result['program_id']=row['_resolved_program_id']
    if row.get('_resolved_direction_id'):result['direction_id']=row['_resolved_direction_id']
    if result.get('program_id'):
        # The validated display name is retained; the stable ID is used by scheduling.
        result['department']=str(row.get('department') or '').strip() or result.get('department')
    result.pop('program_code',None);result.pop('direction_name',None);result.pop('offering_unit_code',None)
    return result
