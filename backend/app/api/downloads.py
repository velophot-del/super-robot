"""管理员数据下载中心：班级、教师、课程、教室与课表。"""
from datetime import datetime
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.class_info import ClassInfo
from app.models.classroom import Classroom
from app.models.course import Course
from app.models.planning import AcademicProgram, OfferingUnit, PublishedPlan
from app.models.schedule import ScheduleEntry
from app.models.teacher import Teacher
from app.services.auth import require_user
from app.services.authorization import check
from app.services.course_catalog import course_schedule_info
from app.services.schedule_versions import resolve_schedule_version

router = APIRouter(prefix='/api/downloads', tags=['数据下载'])

DATASETS = {
    'classes': '班级数据',
    'teachers': '教师数据',
    'courses': '课程数据',
    'classrooms': '教室数据',
    'schedule': '课表数据',
    'all': '全部数据',
}

DAY_NAMES = {1: '周一', 2: '周二', 3: '周三', 4: '周四', 5: '周五', 6: '周六', 7: '周日'}
XLSX_MEDIA = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def _weekdays_text(days):
    return '、'.join(DAY_NAMES[day] for day in days if day in DAY_NAMES)


def _fill_sheet(ws, title, headers, rows):
    ws.title = title[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='4472C4')
        cell.alignment = Alignment(horizontal='center')
    for row in rows:
        ws.append(['' if value is None else value for value in row])
    for index, header in enumerate(headers, 1):
        width = max(12, min(40, len(str(header)) * 2 + 4))
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = 'A2'
    if rows:
        ws.auto_filter.ref = ws.dimensions


def _class_rows(db):
    programs = {row.id: row.name for row in db.query(AcademicProgram)}
    return [[row.id, row.name, row.grade, row.student_count, row.department or '',
             programs.get(row.program_id, ''), row.homeroom_teacher or '',
             '有效' if row.is_active else '停用', row.remarks or '']
            for row in db.query(ClassInfo).order_by(ClassInfo.id)]


CLASS_HEADERS = ['ID', '班级名称', '年级', '学生人数', '所属专业', '标准专业', '班主任', '状态', '备注']


def _teacher_rows(db):
    return [[row.id, row.name, row.teacher_code, row.department or '', row.title or '',
             row.position or '', row.review_major or '', row.email or '', row.phone or '',
             row.max_weekly_hours, row.remarks or '']
            for row in db.query(Teacher).order_by(Teacher.id)]


TEACHER_HEADERS = ['ID', '姓名', '工号', '部门', '职称', '职位', '审核专业方向', '邮箱', '电话',
                   '每周最大课时', '备注']


def _course_rows(db):
    units = {row.id: row.name for row in db.query(OfferingUnit)}
    leaders = {row.id: row.name for row in db.query(Teacher)}
    rows = []
    for row in db.query(Course).order_by(Course.id):
        info = course_schedule_info(row)
        rows.append([row.id, row.name, row.course_code, row.course_type, row.course_nature,
                     row.semester, row.cohort or '', info['teaching_weeks'],
                     _weekdays_text(info['weekdays']), info['total_hours'],
                     row.credits, '是' if row.requires_consecutive else '否',
                     '是' if row.requires_lab else '否', row.department or '',
                     units.get(row.offering_unit_id, ''), leaders.get(row.leader_teacher_id, ''),
                     '活动' if row.is_active else '已移出', row.remarks or ''])
    return rows


COURSE_HEADERS = ['ID', '课程名称', '课程编码', '课程类型', '课程性质', '上课学期', '适用年级',
                  '授课周数', '每周上课日', '总课时', '学分',
                  '需连排', '需实验室', '所属专业', '开课单位', '课程负责人', '状态', '备注']


def _classroom_rows(db):
    return [[row.id, row.name, row.room_code, row.building or '', row.floor, row.capacity,
             row.room_type or '', '有' if row.has_multimedia else '无',
             '可用' if row.is_available else '停用', row.remarks or '']
            for row in db.query(Classroom).order_by(Classroom.id)]


CLASSROOM_HEADERS = ['ID', '教室名称', '编号', '教学楼', '楼层', '容量', '类型', '多媒体', '可用', '备注']

SCHEDULE_PLAN_HEADERS = ['课程名称', '课程编码', '教学组', '年级', '专业', '授课教师', '地点',
                         '起始周', '结束周', '周内节次', '计划学时', '人数']
SCHEDULE_LEGACY_HEADERS = ['课程名称', '教师', '班级', '教室', '周次', '星期', '开始节次',
                           '结束节次', '版本']


def _slots_text(slots):
    parts = []
    for slot in slots or []:
        start = slot.get('start', 1)
        count = slot.get('count', 1)
        phase = '（尾周）' if slot.get('phase') == 'tail' else ''
        parts.append(f"周{slot.get('day')} 第{start}—{start + count - 1}节{phase}")
    return '；'.join(parts)


def _schedule_rows(tasks):
    return [[
        task.get('name', ''), task.get('code', ''), task.get('group', ''),
        task.get('cohort', ''), task.get('major') or task.get('department') or '',
        '、'.join(task.get('teacher_names') or []), task.get('room_name') or '',
        task.get('start_week'), task.get('end_week'), _slots_text(task.get('slots')),
        task.get('required_hours'),
        task.get('actual_student_count') or task.get('student_count') or 0,
    ] for task in tasks]


def _schedule_data(db, plan_id=None):
    """返回 (headers, rows, label)。优先已发布快照，其次当前周期过程稿/空白稿，最后旧版排课。"""
    from app.models.planning import AcademicPlan

    snapshot = None
    plan = None
    if plan_id is not None:
        snapshot = (db.query(PublishedPlan).filter_by(plan_id=plan_id)
                    .order_by(PublishedPlan.id.desc()).first())
        plan = db.get(AcademicPlan, plan_id)
        if not snapshot and not plan:
            raise HTTPException(404, '所选周期不存在')
    else:
        plan = db.query(AcademicPlan).order_by(AcademicPlan.id.desc()).first()
        if plan and plan.state == 'published':
            snapshot = (db.query(PublishedPlan).filter_by(plan_id=plan.id)
                        .order_by(PublishedPlan.id.desc()).first())
        elif not plan:
            snapshot = db.query(PublishedPlan).order_by(PublishedPlan.id.desc()).first()

    if snapshot:
        payload = snapshot.payload or {}
        plan_payload = payload.get('plan') or {}
        return (SCHEDULE_PLAN_HEADERS, _schedule_rows(payload.get('tasks') or []),
                f"已发布课表 · {plan_payload.get('name', plan.name if plan else '')}")
    if plan:
        from app.api.planning import plan_info, tasks_for
        plan_payload = plan_info(plan, db)
        tasks = tasks_for(db, plan)
        names = {row.id: row.name for row in db.query(Teacher)}
        for task in tasks:
            if not task.get('teacher_names'):
                task['teacher_names'] = [names.get(tid, f'教师{tid}')
                                         for tid in (task.get('teacher_ids') or [])]
        label = ('空白稿' if not tasks else '过程稿') + f" · {plan_payload.get('name', '')}"
        return SCHEDULE_PLAN_HEADERS, _schedule_rows(tasks), label
    version = resolve_schedule_version(db, None)
    if not version:
        return SCHEDULE_PLAN_HEADERS, [], '暂无排课数据'
    entries = (db.query(ScheduleEntry).filter(ScheduleEntry.schedule_version == version)
               .order_by(ScheduleEntry.week_number, ScheduleEntry.day_of_week,
                         ScheduleEntry.period_start).all())
    courses = {row.id: row.name for row in db.query(Course)}
    teachers = {row.id: row.name for row in db.query(Teacher)}
    classes = {row.id: row.name for row in db.query(ClassInfo)}
    rooms = {row.id: row.name for row in db.query(Classroom)}
    rows = [[courses.get(entry.course_id, ''), teachers.get(entry.teacher_id, ''),
             classes.get(entry.class_id, ''), rooms.get(entry.classroom_id, ''),
             entry.week_number, DAY_NAMES.get(entry.day_of_week, str(entry.day_of_week)),
             entry.period_start, entry.period_end, version]
            for entry in entries]
    return SCHEDULE_LEGACY_HEADERS, rows, f'旧版排课 · {version}'


def _add_schedule_sheet(wb, db, plan_id=None):
    headers, rows, label = _schedule_data(db, plan_id)
    ws = wb.create_sheet()
    _fill_sheet(ws, '课表数据', headers, rows)
    if not rows:
        ws.append([label])


@router.get('/options')
def download_options(user=Depends(require_user), db: Session = Depends(get_db)):
    check(user, 'admin')
    from app.models.planning import AcademicPlan
    plans = [{'id': plan.id, 'name': plan.name, 'state': plan.state}
             for plan in db.query(AcademicPlan).order_by(AcademicPlan.id.desc())]
    return {'datasets': [{'key': key, 'label': label} for key, label in DATASETS.items()],
            'plans': plans, 'legacy_schedule': resolve_schedule_version(db, None)}


@router.get('/{dataset}')
def download_dataset(dataset: str, plan_id: int | None = Query(None),
                     user=Depends(require_user), db: Session = Depends(get_db)):
    check(user, 'admin')
    if dataset not in DATASETS:
        raise HTTPException(404, '未知的数据类型')
    wb = Workbook()
    wb.remove(wb.active)
    if dataset == 'classes':
        _fill_sheet(wb.create_sheet(), '班级数据', CLASS_HEADERS, _class_rows(db))
    elif dataset == 'teachers':
        _fill_sheet(wb.create_sheet(), '教师数据', TEACHER_HEADERS, _teacher_rows(db))
    elif dataset == 'courses':
        _fill_sheet(wb.create_sheet(), '课程数据', COURSE_HEADERS, _course_rows(db))
    elif dataset == 'classrooms':
        _fill_sheet(wb.create_sheet(), '教室数据', CLASSROOM_HEADERS, _classroom_rows(db))
    elif dataset == 'schedule':
        headers, rows, label = _schedule_data(db, plan_id)
        ws = wb.create_sheet()
        _fill_sheet(ws, '课表数据', headers, rows)
        if not rows:
            ws.append([label])
    else:  # all
        _fill_sheet(wb.create_sheet(), '班级数据', CLASS_HEADERS, _class_rows(db))
        _fill_sheet(wb.create_sheet(), '教师数据', TEACHER_HEADERS, _teacher_rows(db))
        _fill_sheet(wb.create_sheet(), '课程数据', COURSE_HEADERS, _course_rows(db))
        _fill_sheet(wb.create_sheet(), '教室数据', CLASSROOM_HEADERS, _classroom_rows(db))
        _add_schedule_sheet(wb, db, plan_id)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    label = DATASETS[dataset]
    filename = f"{label}_{datetime.now():%Y%m%d}.xlsx"
    return Response(output.getvalue(), media_type=XLSX_MEDIA, headers={
        'Content-Disposition': f"attachment; filename*=UTF-8''{quote(filename)}",
        'Cache-Control': 'no-store',
    })
