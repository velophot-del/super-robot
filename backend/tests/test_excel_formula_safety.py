"""导出 Excel 时用户录入的字符串必须按文本写入，不能被解释为公式。"""
import unittest
from io import BytesIO
from openpyxl import load_workbook
from app.services.export_service import ExportService


class ExcelFormulaSafetyTests(unittest.TestCase):
    def test_formula_like_names_are_written_as_text(self):
        entries = [{
            'course_name': '=HYPERLINK("http://example.invalid","x")',
            'teacher_name': '+cmd|calc',
            'class_name': '@SUM(1,1)',
            'classroom_name': '-1+1',
            'week_number': 1, 'day_of_week': 1,
            'period_start': 1, 'period_end': 2,
            'is_consecutive': True, 'status': 'confirmed',
        }]
        book = load_workbook(BytesIO(ExportService.export_to_excel(
            entries, {'version': '=1+1', 'total_entries': 1})))

        all_rows = book['全部排课数据']
        self.assertEqual(all_rows.cell(2, 2).value, entries[0]['course_name'])
        for column in range(2, 6):
            self.assertEqual(all_rows.cell(2, column).data_type, 's')

        stats = book['统计信息']
        self.assertEqual(stats.cell(3, 2).data_type, 's')
        self.assertEqual(stats.cell(3, 2).value, '=1+1')

        timetable = book['按班级课表']
        self.assertEqual(timetable.cell(3, 2).data_type, 's')
        self.assertIn('=HYPERLINK', str(timetable.cell(3, 2).value))
