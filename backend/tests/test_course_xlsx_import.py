from io import BytesIO
import unittest
from openpyxl import Workbook
from unittest.mock import patch
import test_debug_mode
from app.models.course import Course


class CourseXlsxImportTests(unittest.IsolatedAsyncioTestCase):
    asyncSetUp=test_debug_mode.DebugModeTests.asyncSetUp
    asyncTearDown=test_debug_mode.DebugModeTests.asyncTearDown

    async def test_preview_is_read_only_and_confirmation_imports_only_valid_new_rows_once(self):
        wb=Workbook();one=wb.active;one.title='课程表A'
        one.append(['课程名称','课程编码','上课学期','总课时','课程类型','课程性质'])
        one.append(['构成设计','A101',1,32,'专业课','必修课'])
        one.append(['错误学期','BAD',13,32,'专业课','必修课'])
        two=wb.create_sheet('课程表B');two.append(['课程名称','课程编码','上课学期','总课时'])
        two.append(['重复课程','A101',2,32])
        out=BytesIO();wb.save(out)
        headers={'Authorization':'Bearer '+(await self.client.post('/api/debug/session',json={'role':'admin'})).json()['access_token']}
        preview=await self.client.post('/api/io/xlsx/courses/preview',headers=headers,
            files={'file':('courses.xlsx',out.getvalue())})
        self.assertEqual(preview.status_code,200,preview.text)
        body=preview.json()
        self.assertEqual(self.db.query(Course).count(),0)
        self.assertEqual(body['counts'],{'new':1,'duplicate':1,'invalid':1})
        self.assertIn('课程表B',body['rows'][2]['source_sheet'])
        confirmed=await self.client.post('/api/io/xlsx/courses/confirm',headers=headers,
            json={'preview_id':body['preview_id'],'token':body['token']})
        self.assertEqual(confirmed.status_code,200,confirmed.text)
        self.assertEqual(confirmed.json()['imported'],1)
        self.assertEqual(self.db.query(Course).count(),1)
        replay=await self.client.post('/api/io/xlsx/courses/confirm',headers=headers,
            json={'preview_id':body['preview_id'],'token':body['token']})
        self.assertEqual(replay.status_code,409)

    async def test_preview_accepts_legacy_course_schedule_columns(self):
        wb=Workbook();ws=wb.active;ws.title='2025届课程'
        ws.append(['课程名称','课程编码','上课学期','学期教学次数','每周上课次数','每次课时'])
        ws.append(['版式设计','D2025',2,18,1,3])
        out=BytesIO();wb.save(out)
        headers={'Authorization':'Bearer '+(await self.client.post('/api/debug/session',json={'role':'admin'})).json()['access_token']}
        preview=await self.client.post('/api/io/xlsx/courses/preview',headers=headers,data={'cohort':'2025'},
            files={'file':('2025_courses.xlsx',out.getvalue())})
        self.assertEqual(preview.status_code,200,preview.text)
        row=preview.json()['rows'][0]
        self.assertEqual(row['status'],'new',row)
        self.assertEqual(row['data']['cohort'],'2025')
        self.assertEqual(row['data']['total_hours'],54)
        self.assertEqual(row['data']['teaching_weeks'],18)

    async def test_preview_assigns_code_for_blank_code_and_confirmation_keeps_it(self):
        wb=Workbook();ws=wb.active
        ws.append(['课程名称','课程编码','上课学期','总课时','适用年级','所属专业'])
        ws.append(['构成设计',None,3,32,2024,'视觉传达设计'])
        out=BytesIO();wb.save(out)
        headers={'Authorization':'Bearer '+(await self.client.post('/api/debug/session',json={'role':'admin'})).json()['access_token']}
        preview=await self.client.post('/api/io/xlsx/courses/preview',headers=headers,
            files={'file':('courses.xlsx',out.getvalue())})
        self.assertEqual(preview.status_code,200,preview.text)
        row=preview.json()['rows'][0]
        self.assertEqual(row['course_code'],'24010301')
        confirmed=await self.client.post('/api/io/xlsx/courses/confirm',headers=headers,
            json={'preview_id':preview.json()['preview_id'],'token':preview.json()['token']})
        self.assertEqual(confirmed.status_code,200,confirmed.text)
        self.assertEqual(self.db.query(Course).one().course_code,'24010301')

    async def test_selected_cohort_is_applied_before_auto_code_generation(self):
        wb=Workbook();ws=wb.active
        ws.append(['课程名称','课程编码','上课学期','总课时','所属专业'])
        ws.append(['构成设计',None,3,32,'视觉传达设计'])
        out=BytesIO();wb.save(out)
        headers={'Authorization':'Bearer '+(await self.client.post('/api/debug/session',json={'role':'admin'})).json()['access_token']}
        preview=await self.client.post('/api/io/xlsx/courses/preview',headers=headers,data={'cohort':'2024'},
            files={'file':('2024_courses.xlsx',out.getvalue())})
        self.assertEqual(preview.status_code,200,preview.text)
        row=preview.json()['rows'][0]
        self.assertEqual(row['course_code'],'24010301')
        self.assertEqual(row['data']['cohort'],'2024')
        self.assertEqual(row['status'],'new',row)

    def test_value_reader_retries_when_workbook_has_invalid_fill_style(self):
        from openpyxl import load_workbook as real_load_workbook
        from app.services.spreadsheet_import import _load_workbook_for_values
        wb=Workbook();ws=wb.active
        ws.append(['课程名称','课程编码','上课学期','总课时'])
        ws.append(['构成设计','24010301',3,32])
        out=BytesIO();wb.save(out)
        calls=[0]
        def fail_once_then_load(*args,**kwargs):
            if calls[0]==0:
                calls[0]+=1
                raise TypeError("expected <class 'openpyxl.styles.fills.Fill'>")
            return real_load_workbook(*args,**kwargs)
        with patch('app.services.spreadsheet_import.load_workbook',side_effect=fail_once_then_load) as loader:
            repaired=_load_workbook_for_values(out.getvalue())
        self.assertEqual(repaired.active['A2'].value,'构成设计')
        self.assertEqual(loader.call_count,2)
        repaired.close()


if __name__=='__main__':
    unittest.main()
