from io import BytesIO
import unittest
import test_debug_mode
from openpyxl import Workbook


class CurriculumMaterialsTests(unittest.IsolatedAsyncioTestCase):
    asyncSetUp = test_debug_mode.DebugModeTests.asyncSetUp
    asyncTearDown = test_debug_mode.DebugModeTests.asyncTearDown

    async def test_upload_preview_activate_and_generate_idempotently(self):
        token=(await self.client.post('/api/debug/session',json={'role':'admin'})).json()['access_token']
        h={'Authorization':'Bearer '+token}
        wb=Workbook();ws=wb.active
        ws.append(['课程编码','课程名称','学分','总学时','开课学期','专业方向','领域标签'])
        ws.append(['DES101','Typography',1.5,32,2,'Design','type,layout'])
        data=BytesIO();wb.save(data)
        r=await self.client.post('/api/planning/curricula/import',headers=h,
            data={'major':'Design','cohort':'2026','revision':'1'},
            files={'file':('scheme.xlsx',data.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
        self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(r.json()['rows'][0]['credits'],1.5)
        doc=r.json()['id']
        self.assertEqual((await self.client.post(f'/api/planning/curricula/{doc}/activate',headers=h)).status_code,200)
        r=await self.client.post('/api/planning/plans',headers=h,json={
            'name':'Term','start_date':'2027-02-01','weeks':20,'deadline':'2027-01-31T23:59:59',
            'major':'Design','cohort':'2026','semester':2,'groups':['A','B','C','D']})
        pid=r.json()['id']
        for _ in range(2):
            r=await self.client.post(f'/api/planning/plans/{pid}/generate',headers=h)
            self.assertEqual(r.status_code,200,r.text)
            self.assertEqual(r.json()['task_count'],4)
        state=(await self.client.get(f'/api/planning/plans/{pid}',headers=h)).json()
        self.assertEqual(len(state['tasks']),4)
        self.assertEqual(state['tasks'][0]['credits'],1.5)
        self.assertEqual((await self.client.post(f'/api/planning/plans/{pid}/open',headers=h)).status_code,409)

    async def test_xlsx_title_and_missing_semester_preserve_original(self):
        import hashlib
        token=(await self.client.post('/api/debug/session',json={'role':'admin'})).json()['access_token']
        headers={'Authorization':'Bearer '+token}
        wb=Workbook();ws=wb.active;ws.title='专业课'
        ws.append(['人才培养方案']);ws.append(['课程名称','学分','总课时'])
        ws.append(['版式设计',1.5,32]);output=BytesIO();wb.save(output);content=output.getvalue()
        response=await self.client.post('/api/planning/curricula/import',headers=headers,
            data={'major':'Design','cohort':'2026','revision':'1'},
            files={'file':('scheme.xlsx',content,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
        self.assertEqual(response.status_code,200,response.text)
        doc=response.json();self.assertEqual(doc['sha256'],hashlib.sha256(content).hexdigest())
        self.assertFalse(doc['active']);self.assertFalse(doc['rows'][0]['schedulable'])
        self.assertEqual(doc['rows'][0]['source_sheet'],'专业课')
        source=await self.client.get(f"/api/planning/curricula/{doc['id']}/source",headers=headers)
        self.assertEqual(source.content,content)

    async def test_review_rows_are_preserved_but_not_synced_as_courses(self):
        from app.models.course import Course
        token=(await self.client.post('/api/debug/session',json={'role':'admin'})).json()['access_token']
        headers={'Authorization':'Bearer '+token}
        wb=Workbook();ws=wb.active;ws.title='四年制艺术类'
        ws.append(['课程编码','课程名称','学分','总学时','开课学期'])
        ws.append(['NORMAL1','正常课程',2,32,2])
        ws.append(['','讲座','校级讲座0.5分/场，院级讲座0.25/场','',2])
        ws.append(['说明','','','',''])
        ws.append(['','参考课程',2,'404（参考值）',2])
        ws.append(['','合计','=SUM(C2:C5)','=SUM(D2:D5)',''])
        output=BytesIO();wb.save(output)
        response=await self.client.post('/api/planning/curricula/import',headers=headers,
            data={'major':'Design','cohort':'2026','revision':'1'},
            files={'file':('mixed.xlsx',output.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
        self.assertEqual(response.status_code,200,response.text)
        doc=response.json();self.assertEqual(len(doc['rows']),5)
        result=await self.client.post(f"/api/planning/curricula/{doc['id']}/activate",headers=headers)
        self.assertEqual(result.status_code,200,result.text)
        self.assertEqual(result.json()['course_library']['added'],1)
        self.assertEqual(self.db.query(Course).count(),1)
        self.assertEqual(self.db.query(Course).one().course_code,'NORMAL1')
