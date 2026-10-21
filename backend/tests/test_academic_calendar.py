from io import BytesIO
import unittest
import test_debug_mode
from openpyxl import load_workbook
from app.services.excel_timetable import build_visual_excel


class AcademicCalendarTests(unittest.IsolatedAsyncioTestCase):
    asyncSetUp=test_debug_mode.DebugModeTests.asyncSetUp
    asyncTearDown=test_debug_mode.DebugModeTests.asyncTearDown

    async def test_opening_date_and_independent_monday_persist_and_export(self):
        token=(await self.client.post('/api/debug/session',json={'role':'admin'})).json()['access_token']
        headers={'Authorization':'Bearer '+token}
        payload=dict(name='Leap term',opening_date='2028-02-29',start_date='2028-03-06',weeks=20,deadline='2028-02-20T12:00:00',major='Design')
        result=await self.client.post('/api/planning/plans',headers=headers,json=payload)
        self.assertEqual(result.status_code,200,result.text)
        plan=result.json()
        saved=(await self.client.get(f"/api/planning/plans/{plan['id']}",headers=headers)).json()
        self.assertEqual(saved['opening_date'],'2028-02-29')
        self.assertEqual(saved['start_date'],'2028-03-06')
        sheet=load_workbook(BytesIO(build_visual_excel([{'plan':plan,'tasks':[]}],'Calendar'))).active
        self.assertEqual(sheet.cell(4,4).value,'3.6—3.12')
        self.assertIn('2028-02-29',sheet['A2'].value)
        for changes in [dict(start_date='2028-02-29'),dict(opening_date='2027-02-29')]:
            invalid=await self.client.post('/api/planning/plans',headers=headers,json={**payload,**changes})
            self.assertEqual(invalid.status_code,422)
        legacy={k:v for k,v in payload.items() if k!='opening_date'}
        result=await self.client.post('/api/planning/plans',headers=headers,json=legacy)
        self.assertEqual(result.status_code,200,result.text)
        self.assertIsNone(result.json()['opening_date'])
