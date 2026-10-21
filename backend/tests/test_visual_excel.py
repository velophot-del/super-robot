from datetime import date, datetime
from io import BytesIO
import unittest
import test_debug_mode
from openpyxl import load_workbook
from app.models.planning import AcademicPlan, PublishedPlan
from app.services.excel_timetable import build_visual_excel


def task(name,start=1,end=4):
    return dict(name=name,code=name,group='G',credits=2,required_hours=32,start_week=start,end_week=end,slots=[],teacher_ids=[7],teacher_names=['Teacher'],room_name='Room')


class VisualExcelTests(unittest.TestCase):
    def test_twenty_weeks_parallel_lanes_and_safe_strings(self):
        batches=[dict(plan=dict(start_date='2027-02-01',major='Design',cohort='2026'),tasks=[task('=Danger'),task('Second',2,3),task('Last',19,19)])]
        book=load_workbook(BytesIO(build_visual_excel(batches,'Test')))
        sheet=book.active
        self.assertEqual(sheet.max_column,23)
        self.assertEqual(sheet.cell(3,4).value,'第1周')
        self.assertEqual(sheet.cell(3,23).value,'第20周')
        self.assertIn('D5:G5',[str(r) for r in sheet.merged_cells.ranges])
        self.assertIn('Second',sheet.cell(6,5).value)
        self.assertIn('Last',sheet.cell(5,22).value)
        self.assertEqual(sheet.cell(5,23).value,'考试')
        self.assertEqual(book['课程明细']['A2'].data_type,'s')
        self.assertEqual(book['课程明细'].max_row,4)
        self.assertEqual(sheet.freeze_panes,'D5')

    def test_export_follows_plan_weeks_instead_of_fixed_twenty(self):
        batches=[dict(plan=dict(start_date='2027-02-01',weeks=18,major='Design'),tasks=[])]
        book=load_workbook(BytesIO(build_visual_excel(batches,'Test')))
        sheet=book.active
        self.assertEqual(sheet.title,'18周课表')
        self.assertEqual(sheet.max_column,21)
        self.assertEqual(sheet.cell(3,21).value,'第18周')
        self.assertEqual(sheet.cell(3,20).value,'第17周')


class CollegeExportTests(unittest.IsolatedAsyncioTestCase):
    asyncSetUp=test_debug_mode.DebugModeTests.asyncSetUp
    asyncTearDown=test_debug_mode.DebugModeTests.asyncTearDown

    async def test_all_majors_latest_snapshot_matching_dates_and_duration(self):
        for pid,start,weeks in [(1,'2027-02-01',20),(2,'2027-02-01',20),(3,'2027-02-08',20),(4,'2027-02-01',18),(5,'2027-02-01',20)]:
            self.db.add(AcademicPlan(id=pid,name=str(pid),start_date=date.fromisoformat(start),weeks=weeks,deadline=datetime(2027,1,1),state='draft'))
            self.db.flush()
            if pid!=5:
                self.db.add(PublishedPlan(plan_id=pid,payload={'plan':dict(id=pid,start_date=start,weeks=weeks,major=f'M{pid}'),'tasks':[task(f'Course{pid}')]}));self.db.flush()
        self.db.add(PublishedPlan(plan_id=1,payload={'plan':dict(id=1,start_date='2027-02-01',weeks=20,major='M1'),'tasks':[task('Newest')]}));self.db.commit()
        admin=(await self.client.post('/api/debug/session',json={'role':'admin'})).json()['access_token']
        response=await self.client.get('/api/planning/exports/college/1',headers={'Authorization':'Bearer '+admin})
        self.assertEqual(response.status_code,200,response.text)
        book=load_workbook(BytesIO(response.content))
        names=[r[0] for r in list(book['课程明细'].values)[1:]]
        self.assertEqual(set(names),{'Newest','Course2'})
        self.assertEqual((await self.client.get('/api/planning/exports/college/5',headers={'Authorization':'Bearer '+admin})).status_code,409)
        teacher=(await self.client.post('/api/debug/session',json={'role':'teacher','teacher_id':7})).json()['access_token']
        self.assertEqual((await self.client.get('/api/planning/exports/college/1',headers={'Authorization':'Bearer '+teacher})).status_code,403)
        self.assertEqual((await self.client.get('/api/planning/exports/college/1')).status_code,401)
