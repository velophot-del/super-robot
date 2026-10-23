"""管理员数据下载中心：主数据、全部数据、课表快照与权限。"""
import unittest
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook
import test_debug_mode
from app.models.class_info import ClassInfo
from app.models.classroom import Classroom
from app.models.course import Course
from app.models.planning import AcademicPlan, PublishedPlan
from app.models.user import User
from app.services.auth import account_token, hash_password


class DownloadTests(unittest.IsolatedAsyncioTestCase):
    asyncSetUp = test_debug_mode.DebugModeTests.asyncSetUp
    asyncTearDown = test_debug_mode.DebugModeTests.asyncTearDown

    async def admin_headers(self):
        token = (await self.client.post('/api/debug/session', json={'role': 'admin'})).json()['access_token']
        return {'Authorization': f'Bearer {token}'}

    async def test_admin_can_download_master_data_and_all(self):
        self.db.add_all([
            ClassInfo(name='视传2601', grade='2026', department='Design', student_count=30),
            Course(name='构成设计', course_code='C1', semester=2, teaching_weeks=8,
                   total_hours=64, weekdays=[1, 3]),
            Classroom(name='Room A', room_code='R1', capacity=40),
        ])
        self.db.commit()
        headers = await self.admin_headers()
        for dataset, sheet in [('classes', '班级数据'), ('teachers', '教师数据'),
                               ('courses', '课程数据'), ('classrooms', '教室数据')]:
            response = await self.client.get(f'/api/downloads/{dataset}', headers=headers)
            self.assertEqual(response.status_code, 200, response.text)
            book = load_workbook(BytesIO(response.content))
            self.assertIn(sheet, book.sheetnames)
        all_response = await self.client.get('/api/downloads/all', headers=headers)
        self.assertEqual(all_response.status_code, 200, all_response.text)
        book = load_workbook(BytesIO(all_response.content))
        self.assertEqual(book.sheetnames, ['班级数据', '教师数据', '课程数据', '教室数据', '课表数据'])
        courses = book['课程数据']
        header_row = [cell.value for cell in courses[1]]
        self.assertIn('授课周数', header_row)
        weeks_column = header_row.index('授课周数') + 1
        self.assertEqual(courses.cell(2, weeks_column).value, 8)

    async def test_schedule_download_uses_published_snapshot(self):
        self.db.add(AcademicPlan(id=1, name='P', start_date=date(2027, 2, 1), weeks=20,
                                 deadline=datetime(2027, 1, 1), state='published'))
        self.db.add(PublishedPlan(plan_id=1, payload={
            'plan': {'name': '2026-2027 第二学期'},
            'tasks': [{'name': '构成设计', 'code': 'C1', 'group': '视传2601', 'cohort': '2026',
                       'teacher_names': ['张老师'], 'room_name': 'Room A', 'start_week': 1,
                       'end_week': 8, 'slots': [{'day': 1, 'start': 1, 'count': 2}],
                       'required_hours': 32}]}))
        self.db.commit()
        headers = await self.admin_headers()
        options = (await self.client.get('/api/downloads/options', headers=headers)).json()
        self.assertEqual(options['plans'][0]['id'], 1)
        response = await self.client.get('/api/downloads/schedule', headers=headers, params={'plan_id': 1})
        self.assertEqual(response.status_code, 200, response.text)
        sheet = load_workbook(BytesIO(response.content))['课表数据']
        self.assertEqual(sheet.cell(2, 1).value, '构成设计')
        self.assertEqual((await self.client.get('/api/downloads/schedule', headers=headers,
                                                params={'plan_id': 999})).status_code, 404)

    async def test_schedule_without_data_returns_404_and_non_admin_is_rejected(self):
        headers = await self.admin_headers()
        self.assertEqual((await self.client.get('/api/downloads/schedule', headers=headers)).status_code, 404)
        denied = User(username='dl_denied', email='dl_denied@test.invalid',
                      password_hash=hash_password('Test123'))
        self.db.add(denied)
        self.db.commit()
        denied_headers = {'Authorization': 'Bearer ' + account_token(denied, self.db)}
        self.assertEqual((await self.client.get('/api/downloads/classes', headers=denied_headers)).status_code, 403)
        self.assertEqual((await self.client.get('/api/downloads/options', headers=denied_headers)).status_code, 403)
