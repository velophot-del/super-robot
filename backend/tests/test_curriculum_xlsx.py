from io import BytesIO
import unittest
from openpyxl import Workbook
from app.services.curriculum_materials import extract_rows


def workbook_bytes(wb):
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


class CurriculumXlsxTests(unittest.TestCase):
    def test_multiple_sheets_titles_merged_headers_and_aliases(self):
        wb = Workbook()
        wb.active.title = '说明'
        wb.active.append(['人才培养方案说明'])
        ws = wb.create_sheet('专业课程')
        ws.append(['2026级人才培养方案'])
        ws.merge_cells('A1:F1')
        ws.append([])
        ws.append(['课程代码', '课程\n名称', '学分（学分）', '学时分配', None, '建议修读学期'])
        for column in ['A', 'B', 'C', 'F']:
            ws.merge_cells(f'{column}3:{column}4')
        ws.merge_cells('D3:E3')
        ws.cell(4, 4, '总计')
        ws.cell(4, 5, '理论')
        ws.append(['DES101', '版式设计', 1.5, 32, 8, '第2学期'])
        other = wb.create_sheet('公共课程')
        other.append(['课程名称', '学分', '总课时', '开设学期'])
        other.append(['通识课', 2, 36, 1])
        rows, errors = extract_rows('plan.XLSX', workbook_bytes(wb), 'Design|2026|1')
        self.assertEqual(errors, [])
        self.assertEqual(len(rows), 2)
        row = rows[0]
        self.assertEqual((row['credits'], row['required_hours'], row['semester']), (1.5, 32, 2))
        self.assertEqual((row['source_sheet'], row['source_row']), ('专业课程', 5))
        self.assertEqual(row['raw_fields']['学时分配 / 总计'], '32')
        self.assertEqual(rows[1]['source_sheet'], '公共课程')

    def test_missing_semester_is_preserved_for_review(self):
        wb = Workbook()
        wb.active.append(['课程名称', '学分', '总学时'])
        wb.active.append(['待定课程', 2, 32])
        rows, errors = extract_rows('plan.xlsx', workbook_bytes(wb))
        self.assertEqual(errors, [])
        self.assertEqual(len(rows), 1)
        self.assertIsNone(rows[0]['semester'])
        self.assertFalse(rows[0]['schedulable'])

    def test_formula_row_is_retained_for_review_without_evaluation(self):
        wb = Workbook()
        wb.active.title = '课程表'
        wb.active.append(['课程名称', '学分', '总学时', '开课学期'])
        wb.active.append(['版式设计', 2, '=16*2', 2])
        rows, errors = extract_rows('plan.xlsx', workbook_bytes(wb))
        self.assertEqual(errors, [])
        self.assertEqual(len(rows), 1)
        self.assertFalse(rows[0]['schedulable'])
        self.assertEqual(rows[0]['raw_fields']['总学时'], '=16*2')
        self.assertIn('公式', '；'.join(rows[0]['review_notes']))

    def test_no_course_table_diagnostic_lists_sheets(self):
        wb = Workbook()
        wb.active.title = '毕业要求'
        wb.active.append(['毕业要求', '指标点'])
        rows, errors = extract_rows('plan.xlsx', workbook_bytes(wb))
        self.assertEqual(rows, [])
        self.assertIn('毕业要求', '；'.join(errors))

    def test_unclear_semester_or_missing_hours_are_retained_without_guessing(self):
        wb = Workbook()
        wb.active.append(['课程名称', '学分', '总学时', '开课学期'])
        wb.active.append(['跨学期待核对', 2, 32, '1-2'])
        wb.active.append(['学时待核对', 2, None, 2])
        rows, errors = extract_rows('plan.xlsx', workbook_bytes(wb))
        self.assertEqual(errors, [])
        self.assertEqual(len(rows), 2)
        self.assertIsNone(rows[0]['semester'])
        self.assertEqual(rows[0]['raw_fields']['开课学期'], '1-2')
        self.assertIsNone(rows[1]['required_hours'])
        self.assertTrue(all(not row['schedulable'] and row['review_notes'] for row in rows))

    def test_lecture_reference_summary_and_unnamed_rows_do_not_block_courses(self):
        wb = Workbook();ws = wb.active;ws.title = '四年制艺术类'
        ws.append(['课程编码', '课程名称', '学分', '总学时', '开课学期'])
        ws.append(['C1', '正常课程', 2, 32, 2])
        ws.append(['', '讲座', '校级讲座0.5分/场，院级讲座0.25/场', '', ''])
        ws.append(['分类说明', '', '', '', ''])
        ws.append(['', '参考模块', 10, '404（参考值）', ''])
        ws.append(['', '合计', '=SUM(C2:C5)', '=SUM(D2:D5)', ''])
        ws.append(['', '参考课程', 40, '216（参考值）', 13])
        rows, errors = extract_rows('plan.xlsx', workbook_bytes(wb))
        self.assertEqual(errors, [])
        self.assertEqual(len(rows), 6)
        self.assertTrue(rows[0]['schedulable'])
        self.assertTrue(all(not row['schedulable'] for row in rows[1:]))
        self.assertTrue(all(row['record_type']=='review' for row in rows[1:]))
        self.assertIsNone(rows[1]['credits'])
        self.assertIsNone(rows[3]['required_hours'])
        self.assertEqual(rows[2]['name'], '')
        self.assertEqual(rows[2]['raw_fields']['课程编码'], '分类说明')

    def test_semester_matrix_uses_only_single_explicit_positive_column(self):
        wb=Workbook();ws=wb.active
        ws.append(['课程编号','课程名称','学分','学时','每周学时数',None,None,None,'上课周数'])
        ws.append([None,None,None,'总学时','第一学年',None,'第二学年',None,None])
        ws.append([None,None,None,None,'一','二','三','四',None])
        ws.append(['C1','正常课程',3,48,None,16,0,0,3])
        ws.append(['C2','多学期课程',3,48,4,4,None,None,6])
        ws.append(['C3','说明课程',2,32,'每学期8学时',None,None,None,None])
        rows,errors=extract_rows('matrix.xlsx',workbook_bytes(wb))
        self.assertEqual(errors,[])
        self.assertEqual(len(rows),3)
        self.assertEqual((rows[0]['semester'],rows[0]['weekly_hours']),(2,16))
        self.assertTrue(rows[0]['schedulable'])
        self.assertIn('学期2周学时',rows[0]['raw_fields'])
        self.assertTrue(all(not row['schedulable'] for row in rows[1:]))
