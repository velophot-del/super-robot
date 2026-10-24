import unittest
from io import BytesIO
from openpyxl import Workbook
import test_debug_mode
from app.models.class_info import ClassInfo
from app.models.planning import AcademicDirection, AcademicProgram, CurriculumDocument, OfferingUnit
from app.models.teacher import Teacher
from app.models.course import Course
from app.services.academic_scope import curriculum_semester
from app.database import _migrate_academic_scope_columns, _seed_academic_catalog_detail, _seed_default_offering_units


class AcademicCatalogTests(unittest.IsolatedAsyncioTestCase):
    asyncSetUp=test_debug_mode.DebugModeTests.asyncSetUp
    asyncTearDown=test_debug_mode.DebugModeTests.asyncTearDown

    async def test_legacy_scope_migration_binds_only_exact_major_labels(self):
        self.db.add_all([
            CurriculumDocument(id=1,name='legacy',major='视觉传达设计',cohort='2026',revision=1,
                               original=b'x',sha256='hash',rows=[]),
            ClassInfo(name='视传2601',grade='2026',department='视觉传达设计'),
            ClassInfo(name='合作2601',grade='2026',department='视觉传达设计（合作办学）'),
        ])
        self.db.commit()
        _migrate_academic_scope_columns(self.engine)
        _migrate_academic_scope_columns(self.engine)
        program=self.db.query(AcademicProgram).filter_by(name='视觉传达设计').one()
        self.assertEqual(self.db.get(CurriculumDocument,1).program_id,program.id)
        self.assertEqual(self.db.query(ClassInfo).filter_by(name='视传2601').one().program_id,program.id)
        self.assertIsNone(self.db.query(ClassInfo).filter_by(name='合作2601').one().program_id)
        self.assertEqual(program.duration_semesters,8)

    async def test_curriculum_uses_configured_program_duration(self):
        doc=CurriculumDocument(name='five-year',major='Design',cohort='2025',revision=1,
            original=b'x',sha256='hash',duration_semesters=10,
            rows=[{'semester':10,'schedulable':True}])
        self.assertEqual(curriculum_semester(doc,'2030-03-04',2),10)

    async def test_restart_migration_preserves_immutable_legacy_documents(self):
        from app.services.immutable_records import install_immutable_triggers
        self.db.add(CurriculumDocument(name='archive',major='Design',cohort='2026',revision=1,
            original=b'archive',sha256='archived',rows=[]))
        self.db.commit()
        install_immutable_triggers(self.engine)
        _migrate_academic_scope_columns(self.engine)
        _migrate_academic_scope_columns(self.engine)
        self.assertIsNone(self.db.query(CurriculumDocument).one().program_id)

    async def test_default_studios_are_seeded_once_and_admin_rename_is_persistent(self):
        _seed_default_offering_units(self.engine)
        names=[row.name for row in self.db.query(OfferingUnit).order_by(OfferingUnit.code)]
        self.assertEqual(names,['视觉策略与叙事设计教研室','包装系统与可持续设计教研室',
                                '数字艺术与媒介创新教研室','视觉系统与时尚传播教研室',
                                '智能交互与体验设计教研室'])
        unit=self.db.query(OfferingUnit).filter_by(code='DEPT-01').one()
        unit.name='视觉设计教研室'
        self.db.commit()
        _seed_default_offering_units(self.engine)
        self.assertEqual(self.db.query(OfferingUnit).filter_by(code='DEPT-01').one().name,'视觉设计教研室')

    async def test_workbook_catalog_seeds_detail_fields_directions_and_preserves_studio_edits(self):
        _seed_default_offering_units(self.engine)
        _seed_academic_catalog_detail(self.engine)
        _seed_academic_catalog_detail(self.engine)
        self.assertEqual(self.db.query(AcademicProgram).count(),6)
        self.assertEqual(self.db.query(AcademicDirection).count(),8)
        visual=self.db.query(AcademicProgram).filter_by(code='130502').one()
        self.assertEqual(visual.degree_category,'艺术学')
        self.assertEqual(visual.duration_semesters,8)
        self.assertEqual(self.db.query(AcademicDirection).filter_by(program_id=visual.id).count(),3)
        visual.source_notes=None
        self.db.commit()
        _seed_academic_catalog_detail(self.engine)
        self.db.refresh(visual)
        self.assertIsNone(visual.source_notes)

    async def test_renaming_studio_updates_teacher_department_and_preserves_course_link(self):
        unit=OfferingUnit(code='DEPT-01',name='视觉传达设计')
        teacher=Teacher(name='教师',teacher_code='T-1',department='视觉传达设计')
        self.db.add_all([unit,teacher]);self.db.flush()
        course=Course(name='课程',course_code='C-1',semester=1,total_hours=32,offering_unit_id=unit.id)
        self.db.add(course);self.db.commit()
        token=(await self.client.post('/api/debug/session',json={'role':'admin'})).json()['access_token']
        response=await self.client.put(f'/api/planning/offering-units/{unit.id}',headers={'Authorization':f'Bearer {token}'},
            json={'name':'视觉教研室'})
        self.assertEqual(response.status_code,200,response.text)
        self.db.refresh(teacher);self.db.refresh(course)
        self.assertEqual(teacher.department,'视觉教研室')
        self.assertEqual(course.offering_unit_id,unit.id)

    async def test_admin_can_update_program_duration_and_active_state(self):
        token=(await self.client.post('/api/debug/session',json={'role':'admin'})).json()['access_token']
        headers={'Authorization':f'Bearer {token}'}
        created=await self.client.post('/api/planning/programs',headers=headers,
            json={'code':'P-1','name':'设计专业','duration_semesters':8})
        self.assertEqual(created.status_code,200,created.text)
        updated=await self.client.put(f"/api/planning/programs/{created.json()['id']}",headers=headers,
            json={'duration_semesters':10,'is_active':False})
        self.assertEqual(updated.status_code,200,updated.text)
        self.assertEqual(updated.json()['duration_semesters'],10)
        self.assertFalse(updated.json()['is_active'])

    async def test_admin_can_manage_direction_and_studio_association(self):
        token=(await self.client.post('/api/debug/session',json={'role':'admin'})).json()['access_token']
        headers={'Authorization':f'Bearer {token}'}
        program=AcademicProgram(code='130502',name='视觉传达设计',duration_semesters=8)
        unit=OfferingUnit(code='DEPT-01',name='视觉策略与叙事设计教研室')
        self.db.add_all([program,unit]);self.db.commit()
        created=await self.client.post('/api/planning/directions',headers=headers,json={
            'program_id':program.id,'name':'品牌与视觉策略','entry_type':'专业方向',
            'offering_unit_id':unit.id,'notes':'备注','source_order':1})
        self.assertEqual(created.status_code,200,created.text)
        self.assertEqual(created.json()['offering_unit_name'],unit.name)
        updated=await self.client.put(f"/api/planning/directions/{created.json()['id']}",headers=headers,
            json={'name':'品牌视觉策略','notes':'已核实'})
        self.assertEqual(updated.status_code,200,updated.text)
        self.assertEqual(updated.json()['name'],'品牌视觉策略')

    async def test_catalog_xlsx_preview_confirm_and_repeat_are_idempotent(self):
        wb=Workbook();details=wb.active;details.title='专业方向目录'
        details.append(['山东工艺美术学院专业方向目录'])
        details.append(['序号','专业名称','专业代码','授予学位','学制','专业方向','对应教研室','专业建设级别','备注'])
        details.append([1,'视觉传达设计','130502','艺术学','四年','品牌与视觉策略','视觉策略与叙事设计教研室','国家级一流专业','备注'])
        details.append([2,'视觉传达设计','130502','艺术学','四年','插画与叙事设计','视觉策略与叙事设计教研室','国家级一流专业',''])
        summary=wb.create_sheet('专业汇总')
        summary.append(['序号','专业名称','专业代码','授予学位','学制','方向数','专业建设级别'])
        summary.append([1,'视觉传达设计','130502','艺术学','四年',2,'国家级一流专业'])
        out=BytesIO();wb.save(out)
        headers={'Authorization':'Bearer '+(await self.client.post('/api/debug/session',json={'role':'admin'})).json()['access_token']}
        preview=await self.client.post('/api/io/xlsx/academic-catalog/preview',headers=headers,
            files={'file':('catalog.xlsx',out.getvalue())})
        self.assertEqual(preview.status_code,200,preview.text)
        body=preview.json()
        self.assertEqual(body['counts']['new'],3)
        self.assertEqual(self.db.query(AcademicProgram).count(),0)
        confirmed=await self.client.post('/api/io/xlsx/academic-catalog/confirm',headers=headers,
            json={'preview_id':body['preview_id'],'token':body['token']})
        self.assertEqual(confirmed.status_code,200,confirmed.text)
        self.assertEqual(self.db.query(AcademicProgram).count(),1)
        self.assertEqual(self.db.query(AcademicDirection).count(),2)
        self.assertEqual(self.db.query(OfferingUnit).count(),1)
        replay=await self.client.post('/api/io/xlsx/academic-catalog/confirm',headers=headers,
            json={'preview_id':body['preview_id'],'token':body['token']})
        self.assertEqual(replay.status_code,409)

    async def test_catalog_template_download_is_the_workbook_template(self):
        headers={'Authorization':'Bearer '+(await self.client.post('/api/debug/session',json={'role':'admin'})).json()['access_token']}
        response=await self.client.get('/api/io/xlsx/academic-catalog/template',headers=headers)
        self.assertEqual(response.status_code,200,response.text)
        self.assertTrue(response.content.startswith(b'PK'))
        preview=await self.client.post('/api/io/xlsx/academic-catalog/preview',headers=headers,
            files={'file':('专业方向与专业代码目录模板.xlsx',response.content,
                           'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
        self.assertEqual(preview.status_code,200,preview.text)
        self.assertGreater(preview.json()['counts']['new'],0)


if __name__=='__main__':
    unittest.main()
