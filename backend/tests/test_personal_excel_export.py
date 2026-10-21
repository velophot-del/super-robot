from datetime import date,datetime
from io import BytesIO
import unittest
import test_debug_mode
from openpyxl import load_workbook
from app.models.planning import AcademicPlan,PublishedPlan,AccountGrant,TeacherBinding,TeacherProfile
from app.models.user import User
from app.services.auth import hash_password,account_token


class PersonalExcelExportTests(unittest.IsolatedAsyncioTestCase):
    asyncSetUp=test_debug_mode.DebugModeTests.asyncSetUp
    asyncTearDown=test_debug_mode.DebugModeTests.asyncTearDown

    async def test_selected_cycle_exports_only_own_latest_snapshot_even_for_dual_admin(self):
        for pid in [1,2]:
            self.db.add(AcademicPlan(id=pid,name=f'Period {pid}',start_date=date(2027,2,1),weeks=20,deadline=datetime(2027,1,1),state='published'))
        self.db.flush()
        def task(name,teacher):
            return {'name':name,'code':name,'group':'G','credits':2,'required_hours':32,'start_week':1,'end_week':4,'slots':[], 'teacher_ids':[teacher],'teacher_names':['Demo'],'source':{}}
        self.db.add(PublishedPlan(plan_id=1,payload={'plan':{'id':1,'start_date':'2027-02-01','weeks':20},'tasks':[task('Old',7)]}));self.db.flush()
        self.db.add_all([PublishedPlan(plan_id=1,payload={'plan':{'id':1,'start_date':'2027-02-01','weeks':20},'tasks':[task('Mine',7),task('Other',8)]}),PublishedPlan(plan_id=2,payload={'plan':{'id':2,'start_date':'2027-02-01','weeks':20},'tasks':[task('Different cycle',7)]})])
        user=User(username='dual',email='dual@test.invalid',password_hash=hash_password('Test123'))
        self.db.add(user);self.db.flush()
        self.db.add_all([AccountGrant(user_id=user.id,capability='admin',scope='all',scope_id=''),AccountGrant(user_id=user.id,capability='teaching',scope='self',scope_id=''),TeacherBinding(user_id=user.id,teacher_id=7),TeacherProfile(teacher_id=7,payload={'self_profile':{'direction':'Design','tags':['type'],'skilled_courses':['Design']}})])
        self.db.commit()
        headers={'Authorization':'Bearer '+account_token(user,self.db)}
        response=await self.client.get('/api/my/plans/1/export',headers=headers)
        self.assertEqual(response.status_code,200,response.text)
        self.assertIn('.xlsx',response.headers['content-disposition'])
        book=load_workbook(BytesIO(response.content))
        self.assertEqual(book.active.title,'20周课表')
        self.assertEqual(book.active.cell(3,23).value,'第20周')
        self.assertNotIn('Other',str(list(book.active.values)))
        sheet=book['课程明细']
        self.assertEqual(sheet.max_row,2)
        self.assertEqual(sheet.cell(2,1).value,'Mine')
        second=await self.client.get('/api/my/plans/2/export',headers=headers)
        self.assertEqual(load_workbook(BytesIO(second.content))['课程明细'].cell(2,1).value,'Different cycle')
        admin=(await self.client.post('/api/debug/session',json={'role':'admin'})).json()['access_token']
        self.assertEqual((await self.client.get('/api/my/plans/1/export',headers={'Authorization':'Bearer '+admin})).status_code,403)
        self.assertEqual((await self.client.get('/api/my/plans/1/export')).status_code,401)
