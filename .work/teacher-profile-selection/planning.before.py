"""教师先申报的阶段式计划工作流。所有决定来自事务内最新版本。"""
from contextlib import contextmanager
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.teacher import Teacher
from app.models.classroom import Classroom
from app.models.planning import (AcademicPlan, TeachingTask, PeriodSlot, TeacherProfile,
    PreferenceDraft, PreferenceSubmission, MatchRun, TeachingAssignment, PlanningEvent, PublishedPlan)
from app.schemas.planning import PlanInput, TaskInput, ProfileInput, DraftInput, SubmitInput, AssignmentInput
from app.services.auth import require_user
from app.services.authorization import can, check
from app.services.stage_planning import task_dict, validate_tasks, capacity
from app.services.teacher_course_matching import score_candidate

router = APIRouter(prefix='/api/planning', tags=['教师选课与阶段计划'])


@contextmanager
def write(db):
    db.rollback()
    if db.bind.dialect.name == 'sqlite':
        db.execute(text('BEGIN IMMEDIATE'))
    try:
        yield
        db.commit()
    except Exception:
        db.rollback()
        raise


def event(db, user, action, payload, plan_id=None):
    db.add(PlanningEvent(plan_id=plan_id, actor=f"{'demo' if getattr(user, 'is_demo', False) else 'user'}:{user.id}",
                         action=action, payload=payload))


def get_plan(db, plan_id):
    plan = db.get(AcademicPlan, plan_id)
    if not plan: raise HTTPException(404, '周期不存在')
    return plan


def tasks_for(db, plan):
    return [task_dict(db, t) for t in db.query(TeachingTask).filter_by(plan_id=plan.id).order_by(TeachingTask.id)]


def latest_submissions(db, plan_id):
    latest = {}
    for row in db.query(PreferenceSubmission).filter_by(plan_id=plan_id).order_by(PreferenceSubmission.id):
        latest[row.teacher_id] = row
    return latest


def teacher_access(user):
    teacher_id = getattr(user, 'teacher_id', None)
    if teacher_id is None: raise HTTPException(403, '账号未绑定教师')
    check(user, 'teaching', {'teacher_id': teacher_id})
    return teacher_id


def can_read(user):
    return any(g['capability'] in {'admin', 'review', 'publish', 'teaching'} for g in getattr(user,'grants',[]))


def plan_info(plan):
    return {'id':plan.id, 'name':plan.name, 'start_date':plan.start_date.isoformat(),
            'weeks':plan.weeks, 'deadline':plan.deadline.isoformat()+'Z', 'state':plan.state,
            'revision':plan.revision, 'rule_confirmed':plan.rule_confirmed, **(plan.settings or {})}


def profile_for(db, teacher_id):
    p = db.get(TeacherProfile, teacher_id)
    return p.payload if p else ProfileInput().model_dump()


def assigned_tasks(db, plan, teacher_id, exclude=None):
    chosen=[]
    for assignment in db.query(TeachingAssignment):
        if assignment.task_id == exclude or teacher_id not in assignment.teacher_ids: continue
        task = db.get(TeachingTask, assignment.task_id); other = db.get(AcademicPlan, task.plan_id)
        from datetime import timedelta
        if other.id == plan.id or (other.state == 'published' and other.start_date < plan.start_date+timedelta(weeks=plan.weeks)
                and other.start_date+timedelta(weeks=other.weeks) > plan.start_date): chosen.append((task_dict(db,task),other))
    return chosen


@router.get('/plans')
def list_plans(user=Depends(require_user), db: Session=Depends(get_db)):
    if not can_read(user): raise HTTPException(403, '账号尚未授予工作职能')
    return [plan_info(p) for p in db.query(AcademicPlan).order_by(AcademicPlan.id.desc())
            if p.state != 'draft' or can(user,'admin')]


@router.post('/plans')
def create_plan(data: PlanInput, user=Depends(require_user), db: Session=Depends(get_db)):
    check(user,'admin')
    with write(db):
        deadline=data.deadline
        if deadline.tzinfo is None: deadline=deadline.replace(tzinfo=ZoneInfo('Asia/Shanghai'))
        plan=AcademicPlan(name=data.name,start_date=data.start_date,weeks=data.weeks,
                          deadline=deadline.astimezone(timezone.utc).replace(tzinfo=None),rule_confirmed=data.rule_confirmed,
                          settings={'major':data.major,'cohort':data.cohort,'semester':data.semester,'groups':data.groups})
        db.add(plan); db.flush(); event(db,user,'create_plan',data.model_dump(mode='json'),plan.id)
        return plan_info(plan)


@router.get('/plans/{plan_id}')
def get_state(plan_id: int, user=Depends(require_user), db: Session=Depends(get_db)):
    plan=get_plan(db,plan_id)
    if not can_read(user) or (plan.state=='draft' and not can(user,'admin')): raise HTTPException(403,'计划尚未开放')
    all_tasks=tasks_for(db,plan)
    tasks=[t for t in all_tasks if can(user,'admin') or can(user,'publish') or
           can(user,'teaching',{'teacher_id':getattr(user,'teacher_id',None)}) or can(user,'review',t)]
    names={t.id:t.name for t in db.query(Teacher)}
    for task in tasks: task['teacher_names']=[names.get(t,'') for t in task['teacher_ids']]
    teacher_id=getattr(user,'teacher_id',None)
    draft=db.query(PreferenceDraft).filter_by(plan_id=plan.id,teacher_id=teacher_id).first() if teacher_id else None
    latest=latest_submissions(db,plan.id)
    submitted=latest.get(teacher_id)
    run=db.query(MatchRun).filter_by(plan_id=plan.id).order_by(MatchRun.id.desc()).first()
    candidates=[]
    if run:
        for c in run.payload['candidates']:
            task=next((t for t in tasks if t['id']==c['task_id']),None)
            if task is None:continue
            if c['teacher_id']==teacher_id or can(user,'review',task): candidates.append({**c,'teacher_name':names.get(c['teacher_id'],'')})
    return {**plan_info(plan),'tasks':tasks,
        'draft':{'revision':draft.revision,'task_ids':draft.task_ids,'plan_revision':draft.plan_revision} if draft else {'revision':0,'task_ids':[],'plan_revision':plan.revision},
        'submission':{'id':submitted.id,'task_ids':submitted.task_ids} if submitted else None,
        'submission_count':len(latest) if can(user,'admin') or any(can(user,'review',t) for t in tasks) else None,
        'match':{'id':run.id,'maximum':105,'candidates':candidates,'suggestions':[
            s for s in run.payload['suggestions'] if any(t['id']==s['task_id'] and can(user,'review',t) for t in tasks)]} if run else None}


@router.get('/profiles')
def list_profiles(user=Depends(require_user), db: Session=Depends(get_db)):
    check(user,'admin')
    return [{'id':t.id,'name':t.name,**profile_for(db,t.id)} for t in db.query(Teacher).order_by(Teacher.id)]


@router.put('/profiles/{teacher_id}')
def update_profile(teacher_id: int, data: ProfileInput, user=Depends(require_user), db: Session=Depends(get_db)):
    check(user,'admin')
    with write(db):
        if not db.get(Teacher,teacher_id): raise HTTPException(404,'教师不存在')
        profile=db.get(TeacherProfile,teacher_id)
        if profile: profile.payload=data.model_dump(); profile.revision+=1
        else: db.add(TeacherProfile(teacher_id=teacher_id,payload=data.model_dump(),revision=1))
        event(db,user,'update_profile',{'teacher_id':teacher_id,**data.model_dump()})
        return {'success':True}


@router.post('/plans/{plan_id}/tasks')
def add_task(plan_id: int, data: TaskInput, user=Depends(require_user), db: Session=Depends(get_db)):
    check(user,'admin')
    with write(db):
        plan=get_plan(db,plan_id)
        if plan.state!='draft': raise HTTPException(409,'开放后的计划内容已锁定；请建立新修订')
        tasks=tasks_for(db,plan)
        if any(t['code']==data.code and t['group']==data.group and not (t['end_week']<data.start_week or t['start_week']>data.end_week) for t in tasks):
            raise HTTPException(409,'同组同课程周段重复，请使用多人名额或不重叠的分段')
        payload=data.model_dump(); validate_tasks(db,plan,[*[t for t in tasks if t.get('start_week')],payload])
        slots=payload.pop('slots'); task=TeachingTask(plan_id=plan.id,payload=payload); db.add(task);db.flush()
        db.add_all([PeriodSlot(task_id=task.id,**s) for s in slots]); db.flush()
        plan.revision+=1; event(db,user,'create_task',data.model_dump(),plan.id)
        return task_dict(db,task)


@router.put('/plans/{plan_id}/tasks/{task_id}')
def update_task(plan_id:int,task_id:int,data:TaskInput,user=Depends(require_user),db:Session=Depends(get_db)):
    check(user,'admin')
    with write(db):
        plan=get_plan(db,plan_id);task=db.get(TeachingTask,task_id)
        if plan.state!='draft':raise HTTPException(409,'开放后内容只读，请建立新修订')
        if not task or task.plan_id!=plan.id:raise HTTPException(404,'任务不存在')
        payload=data.model_dump()
        if task.payload.get('source'):
            for key in ['code','name','credits','required_hours']:
                if payload[key]!=task.payload[key]:raise HTTPException(409,'来源课程属性不可覆盖，请派生培养方案修订')
            payload['source']=task.payload['source']
        others=[t for t in tasks_for(db,plan) if t['id']!=task_id and t.get('start_week')]
        validate_tasks(db,plan,[*others,payload])
        slots=payload.pop('slots');task.payload=payload
        db.query(PeriodSlot).filter_by(task_id=task.id).delete()
        db.add_all([PeriodSlot(task_id=task.id,**s) for s in slots]);db.flush()
        plan.revision+=1;event(db,user,'edit_task',{'task_id':task_id,'data':data.model_dump()},plan.id)
        return task_dict(db,task)


@router.post('/plans/{plan_id}/open')
def open_plan(plan_id: int,user=Depends(require_user),db: Session=Depends(get_db)):
    check(user,'admin')
    with write(db):
        plan=get_plan(db,plan_id); tasks=tasks_for(db,plan)
        if plan.state!='draft' or not tasks: raise HTTPException(409,'只有包含完整任务的草稿可以开放')
        if plan.deadline<=datetime.now(timezone.utc).replace(tzinfo=None): raise HTTPException(409,'申报截止时间已过')
        if any(not task.get('room_id') for task in tasks):raise HTTPException(409,'请在开放填报前落实全部上课地点')
        validate_tasks(db,plan,tasks);plan.state='open';event(db,user,'open',{},plan.id)
        return plan_info(plan)


def editable(plan):
    if plan.state!='open': raise HTTPException(409,'当前不处于志愿填报阶段')
    if plan.deadline<=datetime.now(timezone.utc).replace(tzinfo=None): raise HTTPException(409,'填报已截止')


@router.put('/plans/{plan_id}/draft')
def save_draft(plan_id: int,data: DraftInput,user=Depends(require_user),db: Session=Depends(get_db)):
    tid=teacher_access(user)
    with write(db):
        plan=get_plan(db,plan_id);editable(plan)
        if data.plan_revision!=plan.revision:raise HTTPException(409,'计划已变更，请重新核对')
        ids={t.id for t in db.query(TeachingTask).filter_by(plan_id=plan.id)}
        if len(set(data.task_ids))!=len(data.task_ids) or not set(data.task_ids)<=ids:raise HTTPException(422,'课程任务无效或重复')
        draft=db.query(PreferenceDraft).filter_by(plan_id=plan.id,teacher_id=tid).first()
        if data.revision!=(draft.revision if draft else 0):raise HTTPException(409,'草稿已有新版本，请刷新核对')
        if draft: draft.task_ids=data.task_ids;draft.plan_revision=plan.revision;draft.revision+=1
        else: draft=PreferenceDraft(plan_id=plan.id,teacher_id=tid,task_ids=data.task_ids,plan_revision=plan.revision,revision=1);db.add(draft)
        return {'revision':draft.revision,'task_ids':draft.task_ids,'plan_revision':plan.revision}


@router.post('/plans/{plan_id}/submit')
def submit(plan_id:int,data:SubmitInput,user=Depends(require_user),db:Session=Depends(get_db)):
    tid=teacher_access(user)
    with write(db):
        plan=get_plan(db,plan_id);editable(plan)
        draft=db.query(PreferenceDraft).filter_by(plan_id=plan.id,teacher_id=tid).first()
        if not draft or draft.revision!=data.revision or draft.plan_revision!=plan.revision:raise HTTPException(409,'草稿或计划版本已变更')
        minimum=min(3,db.query(TeachingTask).filter_by(plan_id=plan.id).count())
        if len(draft.task_ids)<minimum:raise HTTPException(422,f'本周期至少提交{minimum}项志愿（最多5项）')
        previous=latest_submissions(db,plan.id).get(tid)
        if previous and previous.draft_revision==draft.revision:return {'id':previous.id,'task_ids':previous.task_ids}
        row=PreferenceSubmission(plan_id=plan.id,teacher_id=tid,task_ids=list(draft.task_ids),plan_revision=plan.revision,draft_revision=draft.revision)
        db.add(row);db.flush();event(db,user,'submit',{'teacher_id':tid,'task_ids':draft.task_ids},plan.id)
        return {'id':row.id,'task_ids':row.task_ids}


@router.post('/plans/{plan_id}/match')
def match(plan_id:int,user=Depends(require_user),db:Session=Depends(get_db)):
    with write(db):
        plan=get_plan(db,plan_id);tasks=tasks_for(db,plan)
        if not tasks or not all(can(user,'review',t) for t in tasks):raise HTTPException(403,'运行全计划匹配需要覆盖所有任务的审核授权')
        if plan.state not in {'open','matched'} or not plan.rule_confirmed:raise HTTPException(409,'规则未确认或阶段不允许匹配')
        latest=latest_submissions(db,plan.id)
        if not latest:raise HTTPException(409,'尚无教师提交意愿')
        candidates=[];profiles={};suggestions=[];selected={}
        for tid,submission in latest.items():
            if submission.plan_revision!=plan.revision:raise HTTPException(409,'存在过时申报，请先处理')
            profile=profile_for(db,tid);p=db.get(TeacherProfile,tid);profiles[str(tid)]=p.revision if p else 0
            for rank,task_id in enumerate(submission.task_ids):
                task=next(t for t in tasks if t['id']==task_id)
                candidates.append(score_candidate(db,plan,task,tid,profile,rank,assigned_tasks(db,plan,tid,task_id)))
        for task in sorted(tasks,key=lambda t:(sum(c['task_id']==t['id'] for c in candidates),t['id'])):
            if task['locked']:
                suggestions.append({'task_id':task['id'],'teacher_ids':task['teacher_ids'],'locked':True});continue
            choices=[]
            pool=sorted([c for c in candidates if c['task_id']==task['id'] and c['complete']],
                        key=lambda c:(-(c['normalized_total'] or 0),c['rank'],c['teacher_id']))
            for candidate in pool:
                tid=candidate['teacher_id'];chosen=assigned_tasks(db,plan,tid,task['id'])+selected.get(tid,[])
                reasons,_=capacity(db,plan,task,tid,profile_for(db,tid),chosen)
                if not reasons: choices.append(tid)
                if len(choices)==task['teacher_count']:break
            if len(choices)==task['teacher_count']:
                for tid in choices:selected.setdefault(tid,[]).append((task,plan))
            else:choices=[]
            suggestions.append({'task_id':task['id'],'teacher_ids':choices,'locked':False})
        payload={'plan_revision':plan.revision,'submission_versions':{str(t):r.id for t,r in latest.items()},
                 'profile_versions':profiles,'candidates':candidates,'suggestions':suggestions,'maximum':105}
        row=MatchRun(plan_id=plan.id,payload=payload);db.add(row);db.flush();plan.state='matched'
        event(db,user,'match',{'match_id':row.id},plan.id)
        return {'id':row.id,**payload}


@router.post('/plans/{plan_id}/assign')
def assign(plan_id:int,data:AssignmentInput,user=Depends(require_user),db:Session=Depends(get_db)):
    with write(db):
        plan=get_plan(db,plan_id);task=db.get(TeachingTask,data.task_id);run=db.get(MatchRun,data.match_id)
        if not task or task.plan_id!=plan.id:raise HTTPException(404,'任务不存在')
        details=task_dict(db,task);check(user,'review',details)
        if plan.state!='matched' or not run or run.plan_id!=plan.id or run.payload['plan_revision']!=plan.revision:raise HTTPException(409,'匹配批次过期')
        latest_run=db.query(MatchRun.id).filter_by(plan_id=plan.id).order_by(MatchRun.id.desc()).first()
        if latest_run[0]!=run.id:raise HTTPException(409,'请使用最新匹配批次')
        if len(data.teacher_ids)!=details['teacher_count'] or len(set(data.teacher_ids))!=len(data.teacher_ids):raise HTTPException(422,'教师人数与名额不符或重复')
        if getattr(user,'teacher_id',None) in data.teacher_ids:raise HTTPException(403,'本人申报必须回避，请交其他审核人')
        submissions=latest_submissions(db,plan.id)
        own_submission=submissions.get(getattr(user,'teacher_id',None))
        if own_submission and task.id in own_submission.task_ids:raise HTTPException(403,'本人申报的任务必须回避，请交其他审核人')
        for tid in data.teacher_ids:
            sub=submissions.get(tid);p=db.get(TeacherProfile,tid)
            if not sub or task.id not in sub.task_ids or run.payload['submission_versions'].get(str(tid))!=sub.id:raise HTTPException(409,'只能审核本批次已提交的意愿')
            if run.payload['profile_versions'].get(str(tid))!=(p.revision if p else 0):raise HTTPException(409,'教师画像已更新，请重新匹配')
            reasons,_=capacity(db,plan,details,tid,profile_for(db,tid),assigned_tasks(db,plan,tid,task.id))
            if reasons:raise HTTPException(409,'；'.join(reasons))
        old=db.get(TeachingAssignment,task.id)
        if old:raise HTTPException(409,'此任务已审核锁定，不可覆盖；重跑将保留它')
        db.add(TeachingAssignment(task_id=task.id,teacher_ids=data.teacher_ids,match_id=run.id,locked=True))
        event(db,user,'assign',data.model_dump(),plan.id)
        return {'success':True}


@router.post('/plans/{plan_id}/publish')
def publish(plan_id:int,user=Depends(require_user),db:Session=Depends(get_db)):
    check(user,'publish')
    with write(db):
        plan=get_plan(db,plan_id);tasks=tasks_for(db,plan)
        if plan.state!='matched':raise HTTPException(409,'计划尚未完成匹配审核')
        validate_tasks(db,plan,tasks)
        if not tasks or any(not t['locked'] or len(t['teacher_ids'])!=t['teacher_count'] or not t.get('room_id') for t in tasks):
            raise HTTPException(409,'尚有教师待审核或地点未落实')
        for task in tasks:
            for tid in task['teacher_ids']:
                reasons,_=capacity(db,plan,task,tid,profile_for(db,tid),assigned_tasks(db,plan,tid,task['id']))
                if reasons:raise HTTPException(409,'；'.join(reasons))
        names={t.id:t.name for t in db.query(Teacher)}
        for task in tasks:task['teacher_names']=[names.get(t,'') for t in task['teacher_ids']]
        plan.state='published';snapshot={'plan':plan_info(plan),'tasks':tasks}
        db.add(PublishedPlan(plan_id=plan.id,payload=snapshot));event(db,user,'publish',snapshot,plan.id)
        return {'success':True}


@router.get('/plans/{plan_id}/export')
def export_plan(plan_id:int,user=Depends(require_user),db:Session=Depends(get_db)):
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from fastapi.responses import Response
    plan=get_plan(db,plan_id)
    snapshot=db.query(PublishedPlan).filter_by(plan_id=plan.id).order_by(PublishedPlan.id.desc()).first()
    if not snapshot:raise HTTPException(409,'只能导出已发布快照')
    tasks=snapshot.payload['tasks']
    if not (can(user,'admin') or can(user,'publish')):
        if can(user,'teaching',{'teacher_id':getattr(user,'teacher_id',None)}):
            tasks=[t for t in tasks if user.teacher_id in t['teacher_ids']]
        else:tasks=[t for t in tasks if can(user,'review',t)]
        if not tasks:raise HTTPException(403,'无可导出的本人或授权范围任务')
    wb=Workbook();ws=wb.active;ws.title='专业课阶段课表'
    ws.append(['课程名称','课程编码','教学组','学分','计划学时','起始周','结束周','周内分布','地点','任课教师','来源版本'])
    for t in tasks:
        ws.append([t['name'],t['code'],t['group'],t['credits'],t['required_hours'],t['start_week'],t['end_week'],
            '；'.join(f"周{s['day']} 第{s['start']}—{s['start']+s['count']-1}节" for s in t['slots']),
            t.get('room_name',''), '、'.join(t.get('teacher_names',[])),str(t.get('source',{}))])
    for row in ws:
        for cell in row:
            if isinstance(cell.value,str):cell.data_type='s'
            cell.alignment=Alignment(wrap_text=True,vertical='center')
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='533AFD')
    for column in ws.columns:ws.column_dimensions[column[0].column_letter].width=22
    ws.freeze_panes='A2';ws.page_setup.orientation='landscape';output=BytesIO();wb.save(output)
    return Response(output.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition':f'attachment; filename="plan-{plan.id}-snapshot-{snapshot.id}.xlsx"'})


@router.get('/plans/{plan_id}/my-tasks')
def my_tasks(plan_id:int,user=Depends(require_user),db:Session=Depends(get_db)):
    tid=teacher_access(user);plan=get_plan(db,plan_id)
    return [t for t in tasks_for(db,plan) if tid in t['teacher_ids']]


@router.get('/plans/{plan_id}/events')
def events(plan_id:int,user=Depends(require_user),db:Session=Depends(get_db)):
    check(user,'admin')
    return [{'id':e.id,'actor':e.actor,'action':e.action,'payload':e.payload,'created_at':e.created_at}
            for e in db.query(PlanningEvent).filter_by(plan_id=plan_id).order_by(PlanningEvent.id)]
