import csv
import io
import math
import zipfile
from xml.etree import ElementTree as ET
from fastapi import HTTPException

ALIASES={
    'code':['课程编码','课程代码','code'], 'name':['课程名称','name'],
    'credits':['学分','credits'], 'required_hours':['总学时','计划学时','学时','required_hours'],
    'semester':['开课学期','建议学期','学期','semester'],
    'direction':['专业方向','方向','direction'], 'tags':['领域标签','擅长领域','tags'],
}


def extract_rows(filename,content):
    if len(content)>10*1024*1024:raise HTTPException(413,'文件超过10MiB')
    suffix=filename.lower().rsplit('.',1)[-1]
    tables=[]
    try:
        if suffix=='csv':
            try: decoded=content.decode('utf-8-sig')
            except UnicodeDecodeError:decoded=content.decode('gb18030')
            tables=[list(csv.reader(io.StringIO(decoded)))]
        elif suffix in {'xlsx','docx'}:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist())>50*1024*1024:raise HTTPException(413,'压缩内容展开超过50MiB')
                if suffix=='docx':
                    xml=archive.read('word/document.xml')
                    if b'<!DOCTYPE' in xml:raise HTTPException(422,'不接受含DTD的文档')
                    root=ET.fromstring(xml); ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                    for table in root.findall('.//w:tbl',ns):
                        tables.append([[''.join(cell.itertext()) if not cell.findall('.//w:t',ns) else ''.join(n.text or '' for n in cell.findall('.//w:t',ns))
                                        for cell in row.findall('w:tc',ns)] for row in table.findall('w:tr',ns)])
            if suffix=='xlsx':
                from openpyxl import load_workbook
                wb=load_workbook(io.BytesIO(content),read_only=True,data_only=False)
                tables=[list(ws.iter_rows(values_only=True)) for ws in wb.worksheets];wb.close()
        else:raise HTTPException(422,'支持CSV、XLSX和DOCX表格；PDF及扫描识别未实现')
    except HTTPException:raise
    except Exception as e:raise HTTPException(422,'文件无法解析，请使用规范课程表格') from e
    rows=[];errors=[]
    for table_no,table in enumerate(tables,1):
        header=None
        for line_no,line in enumerate(table,1):
            if len(rows)>10000:raise HTTPException(413,'课程行超过10000')
            values=[str(v).strip() if v is not None else '' for v in line]
            mapping={key:next((i for i,v in enumerate(values) if v in aliases),None) for key,aliases in ALIASES.items()}
            if mapping['code'] is not None and mapping['name'] is not None:
                header=mapping;continue
            if not header or not any(values):continue
            record={k:values[i] if i is not None and i<len(values) else '' for k,i in header.items()}
            try:
                if not record['code'] or not record['name']:raise ValueError('课程编码和名称必填')
                if any(v.startswith('=') for v in values):raise ValueError('课程字段不能使用公式')
                record['credits']=float(record['credits']);record['required_hours']=int(record['required_hours']);record['semester']=int(record['semester'])
                if not math.isfinite(record['credits']) or not 0<=record['credits']<=30 or not 1<=record['required_hours']<=400 or not 1<=record['semester']<=12:raise ValueError('学分、学时或学期超出范围')
                record['tags']=[s.strip() for s in record['tags'].replace('，',',').split(',') if s.strip()]
                rows.append(record)
            except (ValueError,TypeError) as e:errors.append(f'表{table_no}行{line_no}：{e}')
    if not rows:errors.append('没有识别到规范课程表头：课程编码、课程名称、学分、总学时、开课学期')
    if len({r['code'] for r in rows})!=len(rows):errors.append('课程编码重复，须核对后再上传')
    return rows,errors
